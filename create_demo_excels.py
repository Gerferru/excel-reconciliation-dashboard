from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent
OUTPUT_DIR = BASE_DIR / "ejemplos_excel"
OUTPUT_DIR.mkdir(exist_ok=True)

# Estilos visuales para los Excels
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
REGULAR_FONT = Font(name="Calibri", size=10)
BOLD_FONT = Font(name="Calibri", size=10, bold=True)
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")

THIN_BORDER = Border(
    left=Side(style="thin", color="E0E0E0"),
    right=Side(style="thin", color="E0E0E0"),
    top=Side(style="thin", color="E0E0E0"),
    bottom=Side(style="thin", color="E0E0E0"),
)

# Datos de Facturación Interna (Tabla A)
data_facturacion = [
    ("FAC-2026-001", "2026-01-10", "Acero & Metales S.A.", "Suministro de material industrial", 1500.00, "Cobrado"),
    ("FAC-2026-002", "2026-01-12", "Consultora Prisma", "Servicios de auditoría Q1", 3250.75, "Cobrado"),
    ("FAC-2026-003", "2026-01-15", "Logística Rápida", "Transporte de mercancía nacional", 890.00, "Cobrado"),
    ("FAC-2026-004", "2026-01-18", "Diseño Creativo SL", "Desarrollo de identidad visual", 450.00, "Cobrado"),
    ("FAC-2026-005", "2026-01-20", "Hostelería del Norte", "Equipamiento de cocina", 2100.00, "Cobrado"),
    ("FAC-2026-006", "2026-01-22", "Retail Express", "Campañas en punto de venta", 120.50, "Cobrado"),
    ("FAC-2026-007", "2026-01-25", "Editorial Aurora", "Impresión de catálogos 2026", 640.00, "Cobrado"),
    ("FAC-2026-008", "2026-01-27", "Muebles & Espacios", "Mobiliario de oficina", 1800.00, "Cobrado"),
    ("FAC-2026-009", "2026-01-29", "Soluciones Cloud", "Suscripción anual servidor", 4200.00, "Cobrado"),
    ("FAC-2026-010", "2026-01-30", "Talleres Gómez", "Mantenimiento flota de vehículos", 980.20, "Cobrado"),
    ("FAC-2026-011", "2026-02-02", "Agroalimentaria del Sur", "Suministro de materia prima", 3100.00, "Pendiente"),
    ("FAC-2026-012", "2026-02-04", "Farmacia Central", "Lote productos sanitarios", 750.00, "Pendiente"),
    ("FAC-2026-013", "2026-02-05", "Distribuidora Sol", "Envío de pedidos mayoristas", 1420.00, "Pendiente"),
]

# Datos del Extracto Bancario (Tabla B)
# Incluye coincidencias, discrepancias de valor y movimientos bancarios huérfanos
data_banco = [
    ("FAC-2026-001", "2026-01-11", "Acero & Metales S.A.", "Pago Factura FAC-2026-001", 1500.00, "Cobrado"),
    ("FAC-2026-002", "2026-01-13", "Consultora Prisma", "Transferencia recibida FAC-2026-002", 3250.75, "Cobrado"),
    ("FAC-2026-003", "2026-01-16", "Logística Rápida", "Cobro TPV / FAC-2026-003", 890.00, "Cobrado"),
    ("FAC-2026-004", "2026-01-19", "Diseño Creativo SL", "Abono factura FAC-2026-004", 450.00, "Cobrado"),
    ("FAC-2026-005", "2026-01-21", "Hostelería del Norte", "Transferencia FAC-2026-005", 2100.00, "Cobrado"),
    ("FAC-2026-006", "2026-01-23", "Retail Express", "Pago recibido FAC-2026-006", 120.50, "Cobrado"),
    ("FAC-2026-007", "2026-01-26", "Editorial Aurora", "Cobro recibo bancario FAC-2026-007", 640.00, "Cobrado"),
    ("FAC-2026-008", "2026-01-28", "Muebles & Espacios", "Transferencia (-50€ comisión)", 1750.00, "Cobrado"),   # Discrepancia: 1750 vs 1800
    ("FAC-2026-009", "2026-01-30", "Soluciones Cloud", "Recibo devuelto por el banco", 4200.00, "Devuelto"),      # Discrepancia: Devuelto vs Cobrado
    ("FAC-2026-010", "2026-01-31", "Talleres Gómez", "Transferencia neta", 980.00, "Cobrado"),                     # Discrepancia: 980.00 vs 980.20
    ("FAC-2026-014", "2026-02-06", "Ingreso Desconocido", "Transferencia sin referencia clara", 500.00, "Cobrado"),  # Falta en Facturación (Tabla A)
    ("FAC-2026-015", "2026-02-07", "Abono Rectificativo", "Liquidación de intereses extraordinarios", 1150.00, "Cobrado"), # Falta en Facturación (Tabla A)
]


def create_styled_sheet(filename, sheet_title, header_color, headers, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.views.sheetView[0].showGridLines = True

    header_fill = PatternFill(start_color=header_color, end_color=header_color, fill_type="solid")

    # Escribir encabezados
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = header_fill
        cell.alignment = ALIGN_CENTER
        cell.border = THIN_BORDER
        ws.row_dimensions[1].height = 26

    # Escribir filas
    for row_idx, row_data in enumerate(rows, 2):
        ws.row_dimensions[row_idx].height = 20
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = REGULAR_FONT
            cell.border = THIN_BORDER

            if isinstance(value, float):
                cell.number_format = "#,##0.00 €"
                cell.alignment = ALIGN_RIGHT
            elif isinstance(value, str) and (value.startswith("FAC-") or value.startswith("2026-")):
                cell.alignment = ALIGN_CENTER
            else:
                cell.alignment = ALIGN_LEFT

    # Ajustar ancho de columnas
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    file_path = OUTPUT_DIR / filename
    wb.save(file_path)
    print(f"Archivo generado: {file_path}")
    return file_path


def main():
    # Tabla A: Facturación Interna (color azul índigo)
    headers_a = ["ID_Factura", "Fecha_Emision", "Cliente", "Concepto_ERP", "Importe_Total", "Estado"]
    create_styled_sheet(
        "1_Facturacion_Interna.xlsx",
        "Facturacion",
        "4F46E5", # Indigo
        headers_a,
        data_facturacion,
    )

    # Tabla B: Extracto Bancario (color verde esmeralda)
    headers_b = ["ID_Factura", "Fecha_Operacion", "Cliente", "Concepto_Banco", "Importe_Total", "Estado"]
    create_styled_sheet(
        "2_Extracto_Bancario.xlsx",
        "Extracto_Banco",
        "059669", # Emerald
        headers_b,
        data_banco,
    )


if __name__ == "__main__":
    main()
